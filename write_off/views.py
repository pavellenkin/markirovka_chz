from django.shortcuts import render
from django.http import JsonResponse
from django.db import transaction
from django.utils import timezone
from django.views.decorators.cache import cache_control
from django.core.paginator import Paginator
from .models import WriteOffAct, WriteOffCode
import time
from django.db import transaction, OperationalError
import requests
import json
from main.api.unified_authentication_true_api import auth, load_token
import threading
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import re
from datetime import datetime
from django.http import JsonResponse, HttpResponse
import random


@cache_control(no_cache=True)
def write_off(request):
    """Страница списания"""
    return render(request, 'write_off.html')


@cache_control(no_cache=True)
def write_off_check_act(request):
    """Проверка существования акта"""
    if request.method == "POST":
        act_number = request.POST.get("act_number", "").strip()



        if act_number[0] == ']':
            act_number = act_number[3:]

        if not act_number:
            return JsonResponse({"status": "error", "message": "Введите номер акта"})

        try:
            act = WriteOffAct.objects.get(act_number=act_number)
            codes_count = act.codes.count()
            return JsonResponse({
                "status": "exists",
                "act_number": act.act_number,
                "codes_count": codes_count,
                "created_at": act.created_at.strftime("%d.%m.%Y %H:%M"),
                "message": f"Акт №{act_number} найден. Записей: {codes_count}"
            })
        except WriteOffAct.DoesNotExist:
            return JsonResponse({
                "status": "not_exists",
                "act_number": act_number,
                "message": f"Акт №{act_number} не найден. Создать новый?"
            })


@cache_control(no_cache=True)
def write_off_create_act(request):
    """Создание нового акта"""
    if request.method == "POST":
        act_number = request.POST.get("act_number", "").strip()

        if act_number[0] == ']':
            act_number = act_number[3:]

        if not act_number:
            return JsonResponse({"status": "error", "message": "Введите номер акта"})

        try:
            act = WriteOffAct.objects.get(act_number=act_number)
            return JsonResponse({
                "status": "exists",
                "act_number": act.act_number,
                "codes_count": act.codes.count(),
                "message": f"Акт уже существует. Записей: {act.codes.count()}"
            })
        except WriteOffAct.DoesNotExist:
            act = WriteOffAct.objects.create(act_number=act_number)
            return JsonResponse({
                "status": "created",
                "act_number": act.act_number,
                "message": f"Акт №{act_number} создан"
            })


@cache_control(no_cache=True)
def write_off_get_codes(request):
    """Получение списка кодов для акта"""
    if request.method == "POST":
        act_number = request.POST.get("act_number", "").strip()

        if act_number[0] == ']':
            act_number = act_number[3:]

        if not act_number:
            return JsonResponse({"status": "error", "message": "Номер акта не указан"})

        try:
            act = WriteOffAct.objects.get(act_number=act_number)
            codes = act.codes.values_list('code', flat=True)
            return JsonResponse({
                "status": "success",
                "codes": list(codes),
                "count": len(codes)
            })
        except WriteOffAct.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Акт не найден"})


@cache_control(no_cache=True)
def write_off_save_codes(request):
    """Сохранение кодов в акт с обработкой блокировки БД"""
    if request.method == "POST":
        act_number = request.POST.get("act_number", "").strip()

        if act_number and act_number[0] == ']':
            act_number = act_number[3:]

        codes_json = request.POST.get("codes", "[]")
        create_act = request.POST.get("create_act", "false") == "true"

        try:
            codes = json.loads(codes_json)

            rebase_codes = []
            for item_code in codes:
                if item_code and item_code[0] == ']':
                    item_code = item_code[3:]
                if '\x1d' in item_code:
                    item_code = item_code.split('\x1d')[0]
                    print(item_code)
                item_code = item_code.strip()
                if item_code:  # Проверяем, что код не пустой
                    rebase_codes.append(item_code)
            codes = rebase_codes

        except json.JSONDecodeError:
            return JsonResponse({"status": "error", "message": "Неверный формат данных"})

        if not act_number:
            return JsonResponse({"status": "error", "message": "Номер акта не указан"})

        if not codes:
            return JsonResponse({"status": "error", "message": "Нет кодов для сохранения"})

        # === ФУНКЦИЯ СОХРАНЕНИЯ С ПОВТОРНЫМИ ПОПЫТКАМИ ===
        def save_with_retry(max_retries=5):
            last_error = None

            for attempt in range(max_retries):
                try:
                    with transaction.atomic():
                        # Получаем или создаём акт с блокировкой
                        try:
                            act = WriteOffAct.objects.select_for_update().get(act_number=act_number)
                        except WriteOffAct.DoesNotExist:
                            if create_act:
                                act = WriteOffAct.objects.create(act_number=act_number)
                            else:
                                return JsonResponse({"status": "error", "message": "Акт не найден"})

                        new_codes = []
                        overwritten_codes = []

                        for code in codes:
                            code = code.strip()
                            if not code:
                                continue

                            obj, created = WriteOffCode.objects.get_or_create(
                                act=act,
                                code=code,
                                defaults={'is_overwritten': False}
                            )

                            if not created:
                                obj.is_overwritten = True
                                obj.overwritten_at = timezone.now()
                                obj.save()
                                overwritten_codes.append(code)
                            else:
                                new_codes.append(code)

                        print(f"Сохранено кодов: {len(new_codes)} новых, {len(overwritten_codes)} перезаписано")

                        return JsonResponse({
                            "status": "success",
                            "message": f"Сохранено: {len(new_codes)} новых, {len(overwritten_codes)} перезаписано",
                            "new_count": len(new_codes),
                            "overwritten_count": len(overwritten_codes)
                        })

                except OperationalError as e:
                    last_error = e
                    if 'database is locked' in str(e):
                        wait_time = 0.5 * (attempt + 1)
                        print(f"⚠️ База данных заблокирована, попытка {attempt + 1}/{max_retries}, ждём {wait_time}с")
                        time.sleep(wait_time)
                        continue
                    else:
                        # Другая ошибка OperationalError - пробрасываем дальше
                        raise e
                except Exception as e:
                    last_error = e
                    raise e

            # Если все попытки не удались
            return JsonResponse({
                "status": "error",
                "message": f"База данных временно недоступна. Попробуйте позже. {str(last_error) if last_error else ''}"
            })

        return save_with_retry()

@cache_control(no_cache=True)
def write_off_acts_list(request):
    """Страница со списком актов"""
    return render(request, 'write_off_acts_list.html')


@cache_control(no_cache=True)
def write_off_get_acts_data(request):
    """Получение данных для списка актов (AJAX)"""
    search = request.GET.get('search', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    page = int(request.GET.get('page', 1))
    per_page = 20

    queryset = WriteOffAct.objects.prefetch_related('codes').all()

    if search:
        queryset = queryset.filter(act_number__icontains=search)

    if date_from:
        queryset = queryset.filter(created_at__date__gte=date_from)

    if date_to:
        queryset = queryset.filter(created_at__date__lte=date_to)

    paginator = Paginator(queryset, per_page)
    page_obj = paginator.get_page(page)

    data = []
    for act in page_obj:
        data.append({
            'id': act.id,
            'act_number': act.act_number,
            'created_at': act.created_at.strftime("%d.%m.%Y %H:%M"),
            'codes_count': act.codes.count(),
            'overwritten_count': act.codes.filter(is_overwritten=True).count()
        })

    return JsonResponse({
        'status': 'success',
        'data': data,
        'total': paginator.count,
        'page': page,
        'pages': paginator.num_pages,
        'has_next': page_obj.has_next(),
        'has_previous': page_obj.has_previous()
    })


# ============================================ //
# ====== ФУНКЦИЯ СОЗДАНИЯ EXCEL ====== //
# ============================================ //

def create_export_excel(act_number, results):
    """Создание Excel файла с результатами"""
    wb = Workbook()
    sheet_name = f"Акт_{act_number}"
    sheet_name = re.sub(r'[\\/*?:\[\]]', '_', sheet_name)
    if len(sheet_name) > 31:
        sheet_name = sheet_name[:31]
    wb.active.title = sheet_name
    ws = wb.active

    # Стили
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F1F1F", end_color="1F1F1F", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Заголовки
    headers = ['№', 'Код идентификации', 'Наименование', 'Товарная группа', 'ID группы', 'Статус', 'Бренд', 'Владелец']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    # Данные
    for idx, (code, info) in enumerate(results.items(), 1):
        row_num = idx + 1

        # Номер
        cell = ws.cell(row=row_num, column=1, value=idx)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

        # Код
        cell = ws.cell(row=row_num, column=2, value=code)
        cell.alignment = Alignment(horizontal="left", wrap_text=True)
        cell.border = border

        # Наименование
        name = info.get('productName', '')
        if info.get('error'):
            name = f"ОШИБКА: {info['error']}"
            cell = ws.cell(row=row_num, column=3, value=name)
            cell.font = Font(color="B33C3C")
        else:
            cell = ws.cell(row=row_num, column=3, value=name)
        cell.alignment = Alignment(horizontal="left", wrap_text=True)
        cell.border = border

        # Товарная группа
        cell = ws.cell(row=row_num, column=4, value=info.get('productGroup', ''))
        cell.alignment = Alignment(horizontal="left")
        cell.border = border

        # ID группы
        cell = ws.cell(row=row_num, column=5, value=info.get('productGroupId', ''))
        cell.alignment = Alignment(horizontal="left")
        cell.border = border

        # Статус
        status = info.get('status', '')
        if info.get('error'):
            status = info['error']
            cell = ws.cell(row=row_num, column=6, value=status)
            cell.font = Font(color="B33C3C")
        else:
            status_map = {
                'INTRODUCED': 'В обороте',
                'RETIRED': 'Выбыл',
                'APPLIED': 'Нанесён',
                'EMPTY': 'Пустой'
            }
            status = status_map.get(status, status)
            cell = ws.cell(row=row_num, column=6, value=status)
        cell.alignment = Alignment(horizontal="left")
        cell.border = border

        # Бренд
        cell = ws.cell(row=row_num, column=7, value=info.get('brand', ''))
        cell.alignment = Alignment(horizontal="left")
        cell.border = border

        # Владелец
        cell = ws.cell(row=row_num, column=8, value=info.get('ownerName', ''))
        cell.alignment = Alignment(horizontal="left")
        cell.border = border

    # Настройка ширины колонок
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 20
    ws.column_dimensions['H'].width = 30

    return wb

# ============================================ //
# ====== ЭКСПОРТ В EXCEL ====== //
# ============================================ //

# Хранилище для статусов экспорта
export_status = {}


@cache_control(no_cache=True)
def write_off_export_excel(request):
    """Экспорт акта в Excel с подробной информацией из API"""
    if request.method == "POST":
        act_number = request.POST.get("act_number", "").strip()
        export_id = request.POST.get("export_id", "")

        if not act_number:
            return JsonResponse({"status": "error", "message": "Номер акта не указан"})

        # Очищаем номер акта
        if act_number and act_number[0] == ']':
            act_number = act_number[3:]

        try:
            act = WriteOffAct.objects.get(act_number=act_number)
        except WriteOffAct.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Акт не найден"})

        codes = list(act.codes.all().order_by('scanned_at').values_list('code', flat=True))

        if not codes:
            return JsonResponse({"status": "error", "message": "Нет кодов для экспорта"})

        # Очищаем коды
        cleaned_codes = []
        for code in codes:
            if code and code[0] == ']':
                code = code[3:]
            if '\x1d' in code:
                code = code.split('\x1d')[0]
            if code:
                cleaned_codes.append(code.strip())
        codes = cleaned_codes

        if not codes:
            return JsonResponse({"status": "error", "message": "Нет кодов для экспорта после очистки"})

        # ===== ВСЕГДА ИСПОЛЬЗУЕМ АСИНХРОННЫЙ РЕЖИМ =====
        # Если export_id не передан - генерируем его
        if not export_id:
            export_id = 'export_' + str(int(time.time() * 1000)) + '_' + ''.join(
                random.choices('0123456789abcdef', k=6))

        # Запускаем фоновую обработку
        thread = threading.Thread(
            target=process_export_background,
            args=(export_id, act_number, codes)
        )
        thread.daemon = True
        thread.start()

        return JsonResponse({
            "status": "processing",
            "export_id": export_id,
            "message": f"Начинаем обработку {len(codes)} кодов"
        })

def process_export_background(export_id, act_number, codes):
    """Фоновая обработка экспорта"""
    try:
        export_status[export_id] = {
            'status': 'processing',
            'progress': 0,
            'total': len(codes),
            'message': 'Начинаем обработку...'
        }

        # Получаем информацию о кодах
        results = info_ki_batch(codes)

        export_status[export_id]['progress'] = 100
        export_status[export_id]['status'] = 'ready'
        export_status[export_id]['message'] = f'Обработано {len(results)} из {len(codes)} кодов'
        export_status[export_id]['results'] = results
        export_status[export_id]['act_number'] = act_number

    except Exception as e:
        export_status[export_id] = {
            'status': 'error',
            'message': str(e)
        }


def process_export_sync(act_number, codes):
    """Синхронная обработка экспорта"""
    results = info_ki_batch(codes)
    wb = create_export_excel(act_number, results)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    import re
    safe_act_number = re.sub(r'[^a-zA-Z0-9]', '_', act_number)
    filename = f"Act_{safe_act_number}_{timestamp}"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response

@cache_control(no_cache=True)
def write_off_export_status(request):
    """Получение статуса фонового экспорта"""
    if request.method == "GET":
        export_id = request.GET.get("export_id", "")

        if not export_id or export_id not in export_status:
            return JsonResponse({"status": "error", "message": "Экспорт не найден"})

        data = export_status[export_id]

        if data['status'] == 'ready':
            return JsonResponse({
                'status': 'ready',
                'message': data.get('message', 'Готово к скачиванию'),
                'progress': data.get('progress', 100),
                'total': data.get('total', 0),
                'download_url': f"/write-off/download-export/?export_id={export_id}"
            })
        elif data['status'] == 'error':
            return JsonResponse({
                'status': 'error',
                'message': data.get('message', 'Ошибка обработки')
            })
        else:
            return JsonResponse({
                'status': 'processing',
                'progress': data.get('progress', 0),
                'total': data.get('total', 0),
                'message': data.get('message', 'Обработка...')
            })


@cache_control(no_cache=True)
def write_off_download_export(request):
    """Скачивание готового экспорта"""
    if request.method == "GET":
        export_id = request.GET.get("export_id", "")

        if not export_id or export_id not in export_status:
            return JsonResponse({"status": "error", "message": "Экспорт не найден"})

        data = export_status[export_id]

        if data['status'] != 'ready':
            return JsonResponse({"status": "error", "message": "Экспорт ещё не готов"})

        act_number = data.get('act_number', '')
        results = data.get('results', {})

        wb = create_export_excel(act_number, results)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        # Имя файла только латиница
        import re
        safe_act_number = re.sub(r'[^a-zA-Z0-9]', '_', act_number)
        filename = f"Act_{safe_act_number}_{timestamp}"

        # Удаляем из хранилища
        del export_status[export_id]

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        wb.save(response)
        return response

@cache_control(no_cache=True)
def write_off_get_act_codes(request):
    """Получение кодов акта с пагинацией"""
    if request.method == "GET":
        act_number = request.GET.get("act_number", "").strip()
        page = int(request.GET.get("page", 1))
        per_page = 20

        try:
            act = WriteOffAct.objects.get(act_number=act_number)
        except WriteOffAct.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Акт не найден"})

        codes = act.codes.all().order_by('scanned_at')
        total = codes.count()

        paginator = Paginator(codes, per_page)
        page_obj = paginator.get_page(page)

        data = []
        for code in page_obj:
            data.append({
                'id': code.id,
                'code': code.code,
                'is_overwritten': code.is_overwritten,
                'scanned_at': code.scanned_at.strftime("%d.%m.%Y %H:%M:%S")
            })

        return JsonResponse({
            'status': 'success',
            'codes': data,
            'total': total,
            'page': page,
            'pages': paginator.num_pages,
            'has_next': page_obj.has_next(),
            'has_previous': page_obj.has_previous()
        })


@cache_control(no_cache=True)
def write_off_delete_code(request):
    """Удаление кода из акта"""
    if request.method == "POST":
        code_id = request.POST.get("code_id")
        act_number = request.POST.get("act_number", "").strip()

        try:
            code = WriteOffCode.objects.get(id=code_id)
            act = code.act
            if act.act_number != act_number:
                return JsonResponse({"status": "error", "message": "Код не принадлежит этому акту"})
            code.delete()
            return JsonResponse({"status": "success", "message": "Код удалён"})
        except WriteOffCode.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Код не найден"})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)})


@cache_control(no_cache=True)
def write_off_update_act_number(request):
    """Обновление номера акта"""
    if request.method == "POST":
        old_number = request.POST.get("old_number", "").strip()
        new_number = request.POST.get("new_number", "").strip()

        if not old_number or not new_number:
            return JsonResponse({"status": "error", "message": "Некорректные данные"})

        try:
            act = WriteOffAct.objects.get(act_number=old_number)
            if WriteOffAct.objects.filter(act_number=new_number).exclude(id=act.id).exists():
                return JsonResponse({"status": "error", "message": "Акт с таким номером уже существует"})
            act.act_number = new_number
            act.save()
            return JsonResponse({"status": "success", "message": "Номер акта обновлён"})
        except WriteOffAct.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Акт не найден"})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)})


@cache_control(no_cache=True)
def write_off_delete_act(request):
    """Удаление акта со всеми кодами"""
    if request.method == "POST":
        act_number = request.POST.get("act_number", "").strip()

        try:
            act = WriteOffAct.objects.get(act_number=act_number)
            act.delete()
            return JsonResponse({"status": "success", "message": "Акт удалён"})
        except WriteOffAct.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Акт не найден"})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)})


# ============================================ //
# ====== ФУНКЦИЯ ПОЛУЧЕНИЯ ИНФОРМАЦИИ О КОДАХ ====== //
# ============================================ //

def info_ki_batch(codes, max_batch_size=100):
    """
    Получение информации о кодах пачками (не более 1000 за раз)
    Возвращает словарь {code: info}
    """
    if not codes:
        return {}

    results = {}
    batch_size = min(max_batch_size, 1000)

    for i in range(0, len(codes), batch_size):
        batch = codes[i:i + batch_size]
        print(f"Обработка пачки {i // batch_size + 1}, кодов: {len(batch)}")

        status_load, token = load_token()
        if not status_load:
            status, token = auth()
            if status is False:
                for code in batch:
                    results[code] = {'error': 'Ошибка авторизации'}
                continue

        url = "https://markirovka.crpt.ru/api/v3/true-api/cises/info?pg="
        header = {
            "accept": "application/json",
            "Authorization": f"Bearer {token}",
        }

        try:
            send = requests.post(url, headers=header, json=batch, timeout=30)

            if send.status_code == 200:
                response_data = send.json()
                for item in response_data:
                    cis_info = item.get('cisInfo', {})
                    code = cis_info.get('requestedCis') or cis_info.get('cis', '')
                    if code:
                        if 'errorMessage' in item and item['errorMessage']:
                            results[code] = {
                                'error': item['errorMessage'],
                                'errorCode': item.get('errorCode', ''),
                                'productName': '',
                                'productGroup': '',
                                'productGroupId': '',
                                'status': item['errorMessage'],
                                'brand': '',
                                'ownerName': ''
                            }
                        else:
                            results[code] = {
                                'productName': cis_info.get('productName', ''),
                                'productGroup': cis_info.get('productGroup', ''),
                                'productGroupId': cis_info.get('productGroupId', ''),
                                'status': cis_info.get('status', ''),
                                'brand': cis_info.get('brand', ''),
                                'ownerName': cis_info.get('ownerName', ''),
                                'ownerInn': cis_info.get('ownerInn', ''),
                                'gtin': cis_info.get('gtin', ''),
                                'producedDate': cis_info.get('producedDate', ''),
                                'expirationDate': cis_info.get('expirationDate', ''),
                                'error': None
                            }

                for code in batch:
                    if code not in results:
                        results[code] = {
                            'error': 'Код не найден в ответе',
                            'productName': '',
                            'productGroup': '',
                            'productGroupId': '',
                            'status': 'Код не найден',
                            'brand': '',
                            'ownerName': ''
                        }

            elif send.status_code == 401:
                status, token = auth()
                if status:
                    batch_results = info_ki_batch(batch, max_batch_size)
                    results.update(batch_results)
                else:
                    for code in batch:
                        results[code] = {'error': 'Ошибка авторизации'}
            else:
                error_message = "Код идентификации не найден"
                try:
                    error_data = send.json()
                    if error_data and isinstance(error_data, list) and len(error_data) > 0:
                        if 'errorMessage' in error_data[0]:
                            error_message = error_data[0]['errorMessage']
                except:
                    pass
                for code in batch:
                    results[code] = {
                        'error': error_message,
                        'productName': '',
                        'productGroup': '',
                        'productGroupId': '',
                        'status': error_message,
                        'brand': '',
                        'ownerName': ''
                    }

        except requests.exceptions.Timeout:
            for code in batch:
                results[code] = {'error': 'Таймаут сервиса'}
        except requests.exceptions.ConnectionError:
            for code in batch:
                results[code] = {'error': 'Ошибка соединения'}
        except Exception as e:
            for code in batch:
                results[code] = {'error': str(e)}

    return results


